#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


# 按导入模板推导课程材料期望数量，再逐个核对导出xlsx中的详情Sheet。
REQUIRED_HEADERS = [
    "年级",
    "课程名称",
    "材料名称",
    "出库数量",
    "角色",
    "出库类别",
    "每组数量",
    "每组学生人数",
]


def text(value):
    if value is None:
        return ""
    return str(value).strip()


def number(value):
    if value in (None, ""):
        return 0.0
    try:
        return float(value)
    except Exception:
        return 0.0


def find_source_sheet(workbook):
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        headers = [text(cell.value) for cell in sheet[1]]
        if all(name in headers for name in REQUIRED_HEADERS):
            return sheet
    raise RuntimeError("未找到包含课程出库导入表头的Sheet")


def build_expected_from_template(template_path, default_students, default_teachers):
    workbook = load_workbook(template_path, data_only=True)
    sheet = find_source_sheet(workbook)
    headers = [text(cell.value) for cell in sheet[1]]
    index = {name: headers.index(name) for name in REQUIRED_HEADERS}

    expected = defaultdict(lambda: defaultdict(float))
    for row in sheet.iter_rows(min_row=2, values_only=True):
        grade = text(row[index["年级"]])
        course_name = text(row[index["课程名称"]])
        material_name = text(row[index["材料名称"]])
        role = text(row[index["角色"]])
        outbound_category = text(row[index["出库类别"]]) or "按人"
        outbound_qty = number(row[index["出库数量"]])
        each_group_qty = number(row[index["每组数量"]])
        each_group_student = number(row[index["每组学生人数"]])

        if not grade or not course_name or not material_name:
            continue
        if role not in ("学生", "老师"):
            continue

        if outbound_category == "按组":
            if each_group_qty <= 0 or each_group_student <= 0:
                continue
            group_count = max(1, math.ceil(default_students / each_group_student))
            final_qty = group_count * each_group_qty
        else:
            people = default_students if role == "学生" else default_teachers
            final_qty = outbound_qty * people

        if final_qty <= 0:
            continue

        key = (grade, course_name, role)
        expected[key][material_name] += final_qty

    return expected


def parse_detail_sheet(sheet):
    # 当前详情Sheet格式：
    # A1 返回链接
    # A2 标题
    # A3 年级 / 课程名称
    # 后续仅按学生/老师两个区块输出四列清单。
    grade = text(sheet["B3"].value)
    course_name = text(sheet["D3"].value)
    if not grade or not course_name:
        raise RuntimeError(f"{sheet.title} 缺少年级或课程名称")

    actual = {"学生": [], "老师": []}
    subtotals = {"学生": None, "老师": None}
    total_value = None
    section_rows = {"学生": None, "老师": None}
    subtotal_rows = {"学生": None, "老师": None}

    current_role = None
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_col=4, values_only=True),
        start=1,
    ):
        col_a = text(row[0])
        col_b = text(row[1])
        col_c = text(row[2])
        col_d = text(row[3])

        if col_a == "学生材料清单":
            current_role = "学生"
            section_rows["学生"] = row_index
            continue
        if col_a == "老师材料清单":
            current_role = "老师"
            section_rows["老师"] = row_index
            continue
        if col_a == "材料名称":
            continue
        if col_a == "学生材料小计":
            subtotals["学生"] = number(col_d)
            subtotal_rows["学生"] = row_index
            current_role = None
            continue
        if col_a == "老师材料小计":
            subtotals["老师"] = number(col_d)
            subtotal_rows["老师"] = row_index
            current_role = None
            continue
        if col_a == "合计":
            total_value = number(col_d)
            current_role = None
            continue

        # 跳过标题、信息行、空行。
        if col_a in ("← 返回成本计算", "课程材料成本清单", "年级", ""):
            continue

        if current_role in ("学生", "老师"):
            actual[current_role].append(
                {
                    "material_name": col_a,
                    "unit_price": number(col_b),
                    "quantity": number(col_c),
                    "amount": number(col_d),
                    "row_index": row_index,
                }
            )

    return {
        "grade": grade,
        "course_name": course_name,
        "actual": actual,
        "subtotals": subtotals,
        "subtotal_rows": subtotal_rows,
        "section_rows": section_rows,
        "total": total_value,
    }


def scan_fills(sheet):
    # 详情Sheet要求无背景色，这里抓出所有存在显式填充的单元格。
    bad_cells = []
    for row in sheet.iter_rows():
        for cell in row:
            fill = cell.fill
            pattern = getattr(fill, "patternType", None)
            if pattern and pattern != "none":
                bad_cells.append(cell.coordinate)
    return bad_cells


def compare_sheet(parsed, expected, sheet):
    errors = []
    grade = parsed["grade"]
    course_name = parsed["course_name"]

    for role in ("学生", "老师"):
        expected_items = expected.get((grade, course_name, role), {})
        actual_items = {item["material_name"]: item["quantity"] for item in parsed["actual"][role]}

        if set(expected_items.keys()) != set(actual_items.keys()):
            errors.append(
                f"{sheet.title} {role}材料名称不匹配，期望={sorted(expected_items.keys())}，实际={sorted(actual_items.keys())}"
            )

        for material_name, expected_qty in expected_items.items():
            actual_qty = actual_items.get(material_name)
            if actual_qty is None:
                continue
            if abs(actual_qty - expected_qty) > 0.01:
                errors.append(
                    f"{sheet.title} {role}-{material_name} 数量不匹配，期望={expected_qty:.2f}，实际={actual_qty:.2f}"
                )

        if not expected_items and parsed["actual"][role]:
            errors.append(f"{sheet.title} {role}区块本应为空，但实际存在数据")
        if expected_items and not parsed["actual"][role]:
            errors.append(f"{sheet.title} {role}区块缺少数据")

        line_sum = sum(item["amount"] for item in parsed["actual"][role])
        subtotal = parsed["subtotals"][role]
        expected_subtotal = 0.0 if not expected_items else None
        if subtotal is None:
            if expected_items:
                errors.append(f"{sheet.title} 缺少{role}材料小计")
        else:
            if abs(subtotal - line_sum) > 0.01:
                errors.append(
                    f"{sheet.title} {role}材料小计不等于明细求和，小计={subtotal:.2f}，明细和={line_sum:.2f}"
                )
            if expected_subtotal is not None and abs(subtotal - expected_subtotal) > 0.01:
                errors.append(f"{sheet.title} {role}区块应为空，小计应为0，实际={subtotal:.2f}")

        for item in parsed["actual"][role]:
            display_product = round(item["unit_price"] * item["quantity"], 2)
            if abs(display_product - item["amount"]) > 0.05:
                errors.append(
                    f"{sheet.title} 第{item['row_index']}行金额异常，单价×数量={display_product:.2f}，实际总价={item['amount']:.2f}"
                )

    student_row = parsed["subtotal_rows"]["学生"]
    teacher_row = parsed["section_rows"]["老师"]
    if student_row and teacher_row:
        gap_count = teacher_row - student_row - 1
        if gap_count != 2:
            errors.append(f"{sheet.title} 学生小计与老师清单之间应空2行，实际空{gap_count}行")

    student_subtotal = parsed["subtotals"]["学生"] or 0.0
    teacher_subtotal = parsed["subtotals"]["老师"] or 0.0
    total_value = parsed["total"]
    if total_value is None:
        errors.append(f"{sheet.title} 缺少合计行")
    else:
        expected_total = student_subtotal + teacher_subtotal
        if abs(total_value - expected_total) > 0.01:
            errors.append(
                f"{sheet.title} 合计不正确，期望={expected_total:.2f}，实际={total_value:.2f}"
            )

    filled_cells = scan_fills(sheet)
    if filled_cells:
        preview = ", ".join(filled_cells[:10])
        more = "" if len(filled_cells) <= 10 else f" ... 共{len(filled_cells)}个"
        errors.append(f"{sheet.title} 存在背景色单元格：{preview}{more}")

    return errors


def main():
    parser = argparse.ArgumentParser(description="校验课程出库导出详情Sheet")
    parser.add_argument("--template", required=True, help="导入模板xlsx路径")
    parser.add_argument("--export", required=True, help="导出的xlsx路径")
    parser.add_argument("--students", type=int, default=11, help="默认学生人数")
    parser.add_argument("--teachers", type=int, default=2, help="默认老师人数")
    args = parser.parse_args()

    template_path = Path(args.template)
    export_path = Path(args.export)
    expected = build_expected_from_template(template_path, args.students, args.teachers)

    export_wb = load_workbook(export_path, data_only=True)
    detail_sheets = [
        sheet
        for sheet in export_wb.worksheets
        if text(sheet["A2"].value) == "课程材料成本清单"
    ]
    if not detail_sheets:
        raise RuntimeError("导出文件中未找到课程材料成本清单Sheet")

    all_errors = []
    for sheet in detail_sheets:
        parsed = parse_detail_sheet(sheet)
        all_errors.extend(compare_sheet(parsed, expected, sheet))

    if all_errors:
        print("校验失败：")
        for item in all_errors:
            print(f"- {item}")
        raise SystemExit(1)

    print(f"校验通过：共检查 {len(detail_sheets)} 个课程详情Sheet")


if __name__ == "__main__":
    main()
